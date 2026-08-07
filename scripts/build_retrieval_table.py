"""
Compute the retrieval-curation table at |F|=50 with SimNPO-gentle unlearning.

Reads `cleanslate experiment sheet (1).xlsx`:
  - Pre Eval sheet → baselines per model
  - target-50 sheet → post-unlearning results per (model, curation, source)

For each (curation × source × model) it emits:
  F = (pre_mem_forget - post_mem_forget) / pre_mem_forget * 100   ↑ better
  R = (pre_mem_retain - post_mem_retain) / pre_mem_retain * 100   ↓ better
  QA Δ           = (post_QA_total - pre_QA_total) * 100  pp
  Validation Δ   = (post - pre) * 100  pp  per benchmark

Prints LaTeX rows for the paper table.  Empty / not-yet-run cells become "--".
"""
import sys
from pathlib import Path

import openpyxl

XLSX = Path("/home/arpan/CleanSlate/cleanslate experiment sheet (1).xlsx")

# ── target-50 columns (1-indexed) ──────────────────────────────────────────────
T50_COL = {
    "model": 3, "curation": 4, "source": 5, "method": 6, "setting": 7,
    "preMemF_postMem": 14, "postMemR": 18, "QA": 22,
    "arc": 23, "bbh": 24, "coqa": 25, "drop": 26,
    "gsm8k": 27, "hswag": 29, "heval": 30, "lmbda": 31,
    "mmlu": 32, "triviaqa": 33, "winog": 34,
}
# ── Pre Eval columns (1-indexed) ───────────────────────────────────────────────
PRE_COL = {
    "model": 3, "preMemF": 8, "preMemR": 12, "QA": 16,
    "arc": 17, "bbh": 18, "coqa": 19, "drop": 20,
    "gsm8k": 21, "hswag": 23, "heval": 24, "lmbda": 25,
    "mmlu": 26, "triviaqa": 27, "winog": 28,
}

# Model alias used in the spreadsheet → display name in the LaTeX table
MODELS = [
    ("llama",     "Llama-3.1-8B"),
    ("olmo7b",    "Olmo-3-7B"),
    ("qwen3 8b",  "Qwen3-8B"),
    ("nemotron",  "Nemotron-9B"),
    ("gemma",     "Gemma-3-12B"),
    ("olmo32",    "Olmo-3-32B"),  # also accepts "olmo32b"
]

METHODS = [
    ("BM25 mid",       "BM25", "Midtraining"),
    ("Infinigram mid", "IGM",  "Midtraining"),
    ("BM25 pre",       "BM25", "Pretraining"),
]

# Validation benchmarks shown in the table (in the requested column order)
BENCHMARKS = ["gsm8k", "bbh", "hswag", "winog", "coqa", "heval", "lmbda"]


def load_pre_baselines(wb: openpyxl.Workbook) -> dict[str, dict[str, float]]:
    ws = wb["Pre Eval"]
    out: dict[str, dict[str, float]] = {}
    # Pre-Eval target-50 baselines live in rows 5-10
    for r in range(5, 11):
        cells = [c.value for c in ws[r]]
        model = cells[PRE_COL["model"] - 1]
        if model is None:
            continue
        out[str(model).strip()] = {
            k: cells[PRE_COL[k] - 1]
            for k in PRE_COL
            if k != "model"
        }
    return out


def load_post_results(
    wb: openpyxl.Workbook,
) -> dict[tuple[str, str, str], dict[str, float | None]]:
    """key = (curation, source, model) for simnpo-gentle rows."""
    ws = wb["target-50"]
    out: dict[tuple[str, str, str], dict[str, float | None]] = {}
    for r in range(5, ws.max_row + 1):
        cells = [c.value for c in ws[r]]
        model = cells[T50_COL["model"] - 1]
        curation = cells[T50_COL["curation"] - 1]
        source = cells[T50_COL["source"] - 1]
        method = cells[T50_COL["method"] - 1]
        setting = cells[T50_COL["setting"] - 1]
        if model is None or method != "simnpo" or setting != "gentle":
            continue
        # Skip oracle (curation=Oracle, source="-"): not in this table
        if curation not in {"BM25", "IGM"}:
            continue
        post_mem_f = cells[T50_COL["preMemF_postMem"] - 1]
        # Empty cells → not-yet-run
        if post_mem_f is None:
            continue
        out[(curation, source, str(model).strip())] = {
            "postMemF": post_mem_f,
            "postMemR": cells[T50_COL["postMemR"] - 1],
            "QA": cells[T50_COL["QA"] - 1],
            **{
                k: cells[T50_COL[k] - 1]
                for k in ("arc", "bbh", "coqa", "drop", "gsm8k",
                          "hswag", "heval", "lmbda", "mmlu", "triviaqa", "winog")
            },
        }
    return out


def model_key(alias: str, post: dict[tuple[str, str, str], dict]) -> str | None:
    """Map our display alias to whatever key exists in the post dict."""
    if alias == "olmo32":
        for cand in ("olmo32", "olmo32b"):
            if any(k[2] == cand for k in post):
                return cand
    return alias if any(k[2] == alias for k in post) else None


def compute_row(
    pre: dict[str, float],
    post: dict[str, float | None],
) -> dict[str, str]:
    """Return formatted strings (F, R, QAΔ, and per-benchmark Δ)."""
    cells: dict[str, str] = {}

    # F: forget unmem %  (positive when unlearning works)
    if pre["preMemF"] and post["postMemF"] is not None:
        f_pct = (pre["preMemF"] - post["postMemF"]) / pre["preMemF"] * 100
        cells["F"] = f"{f_pct:.1f}"
    else:
        cells["F"] = "--"

    # R: retain unmem %  (positive = collateral damage)
    if pre["preMemR"] and post["postMemR"] is not None:
        r_pct = (pre["preMemR"] - post["postMemR"]) / pre["preMemR"] * 100
        cells["R"] = f"{r_pct:.1f}"
    else:
        cells["R"] = "--"

    # QA Δ in pp
    if pre["QA"] is not None and post["QA"] is not None:
        d = (post["QA"] - pre["QA"]) * 100
        cells["QA"] = f"{d:+.1f}"
    else:
        cells["QA"] = "--"

    for b in BENCHMARKS:
        if pre.get(b) is not None and post.get(b) is not None:
            d = (post[b] - pre[b]) * 100
            cells[b] = f"{d:+.1f}"
        else:
            cells[b] = "--"

    return cells


def empty_row() -> dict[str, str]:
    return {k: "--" for k in ("F", "R", "QA", *BENCHMARKS)}


def emit_latex(
    pre_b: dict[str, dict[str, float]],
    post_r: dict[tuple[str, str, str], dict],
) -> None:
    """Print the body of the LaTeX tabular (rows only — no preamble)."""
    blocks: list[str] = []
    for method_label, curation, source in METHODS:
        block_lines = [f"\\multirow{{6}}{{*}}{{{method_label}}}"]
        for i, (alias, display) in enumerate(MODELS):
            mk = model_key(alias, post_r)
            key = (curation, source, mk) if mk else None
            if key in post_r and alias in pre_b:
                cells = compute_row(pre_b[alias], post_r[key])
            elif mk and mk in pre_b and key in post_r:
                cells = compute_row(pre_b[mk], post_r[key])
            else:
                cells = empty_row()
            prefix = "& " if i == 0 else "& "
            row = f"{prefix}{display:<14}"
            row += f" & {cells['F']:>5} & {cells['R']:>5}"
            row += f" & {cells['QA']:>5}"
            for b in BENCHMARKS:
                row += f" & {cells[b]:>5}"
            row += " \\\\"
            block_lines.append(row)
        blocks.append("\n".join(block_lines))

    print("\n\\midrule\n".join(blocks))


def main() -> int:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    pre_b = load_pre_baselines(wb)
    post_r = load_post_results(wb)
    emit_latex(pre_b, post_r)
    return 0


if __name__ == "__main__":
    sys.exit(main())
